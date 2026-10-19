"""Step 11: plain CSV lists that go with the map (same registry the site and PDFs use)."""
import csv, json, os
from common import *
def run():
    R = json.load(open(f'{DATA}/registry.json')); L = json.load(open(f'{DATA}/lots_fit.json')); out = os.path.join(REPO, 'h-number-map')
    PV = json.load(open(f'{DATA}/provisional.json')) if os.path.exists(f'{DATA}/provisional.json') else dict(items=[], expected=[], blocks=[])
    pvm = {i['ri']: i for i in PV['items']}
    inferred = [(ri, pvm[ri]) for ri in sorted(pvm) if pvm[ri]['status'] == 'provisional']
    open_items = [(ri, r) for ri, r in enumerate(R['review']) if pvm.get(ri, {}).get('status') != 'provisional']
    taken = {i['prop']['hnum'] for _, i in inferred}
    unplaced = [u for u in R['known_unplaced'] if u['hnum'] not in taken]
    expd = {(e['section'], e['label']): e for e in PV['expected']}
    with open(f'{out}/lots_with_h_numbers.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['h_number', 'section', 'lot', 'map_x', 'map_y', 'outline', 'section_evidence', 'status', 'confidence'])
        for o in sorted(L, key=lambda o: (o['hn'], o['h'])):
            w.writerow([o['h'], o['s'], o['l'], o['cx'], o['cy'], 'drawn' if o['poly'] else 'label only', 'two independent readings' if 'old_ocr_agrees' in o['ev'] else 'inferred from surrounding lots', 'verified', ''])
        for ri, it in inferred:
            r = R['review'][ri]; p = it['prop']
            w.writerow([p['htext'], p['section'], p['label'], r['pt'][0], r['pt'][1], 'drawn' if r['geom'] == 'outlined' else 'label only', 'inferred from the path: ' + it['reason'], 'PROVISIONAL, please confirm', it['conf']])
    with open(f'{out}/unresolved_locations.csv', 'w', newline='') as f:
        w = csv.writer(f); w.writerow(['kind', 'reason', 'h_number_candidate', 'section', 'lot', 'label_on_drawing', 'map_x', 'map_y', 'outline', 'what_we_need_to_know'])
        for ri, r in open_items:
            it = pvm.get(ri)
            w.writerow(['visible lot, identity unresolved', it['reason'] if it else ' | '.join(r.get('reasons', [])), r.get('htext', ''), r.get('section', ''), r.get('legacy', ''), ''.join(r['tokens']), r['pt'][0], r['pt'][1], r['geom'], it['check'] if it else ''])
        for u in unplaced:
            e = expd.get((u['section'], u['legacy']))
            ex = e and e['x'] is not None and e['conf']
            w.writerow(['known H-number, expected spot (no outline drawn)' if ex else 'known H-number, map position unresolved', (e['reason'] if e else '') or u['note'] or 'No label on the drawing resolves to this decoder lot', u['htext'], u['section'], u['legacy'], '', e['x'] if ex else '', e['y'] if ex else '', '', e['check'] if ex else ''])
    # Excel workbook (same registry), shipped next to the PDFs so the site can offer it
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = 'Read me'
    n_ok = sum(1 for o in L if 'old_ocr_agrees' in o['ev']); n_inf = len(L) - n_ok
    rows = [['Cedar Grove Cemetery - H-map lot lists'], [''], ['Lots with H-numbers', len(L), 'Outline (or label position) and a decoded H-number. Section corroborated by two independent readings (%d) or inferred from surrounding lots (%d).' % (n_ok, n_inf)],
            ['Inferred lots (please confirm)', len(inferred), 'Worked out from the path: a lot must stand next to the lots with the H-numbers just before and after it. NOT verified. See the sheet Inferred lots.'],
            ['Lots still open', len(open_items), 'Visible on the drawing; identity could not be tied to one H-number. Not drawn as H-numbers.'],
            ['No map position', len(unplaced), 'Known to the decoder; no lot has been found for them. Some have an expected spot.'],
            [''], ['Section and lot are the cemetery section letters and lot labels. Burial records are not part of this workbook.']]
    for r in rows: ws.append(r)
    ws['A1'].font = Font(bold=True, size=14); ws.column_dimensions['A'].width = 18; ws.column_dimensions['C'].width = 110
    def sheet(name, head, data):
        s = wb.create_sheet(name); s.append(head)
        for c in s[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='0C2340'); c.alignment = Alignment(vertical='center')
        for d in data: s.append(d)
        for i in range(1, len(head) + 1): s.column_dimensions[get_column_letter(i)].width = max(14, min(60, max(len(str(c.value or '')) for c in s[get_column_letter(i)]) + 2))
        s.freeze_panes = 'A2'; s.auto_filter.ref = s.dimensions
    sheet('Lots with H-numbers', ['H-number', 'Section', 'Lot', 'Section evidence', 'Outline'], [[o['h'], o['s'], o['l'], 'two independent readings' if 'old_ocr_agrees' in o['ev'] else 'inferred from surrounding lots', 'drawn' if o['poly'] else 'label only'] for o in sorted(L, key=lambda o: (o['hn'], o['h']))])
    sheet('Inferred lots', ['Status', 'Confidence', 'Proposed H-number', 'Section', 'Lot', 'Label on drawing', 'Why', 'Please check', 'Map x', 'Map y', 'Your answer (Y/N)'], [['PROVISIONAL', i['conf'], i['prop']['htext'], i['prop']['section'], i['prop']['label'], ''.join(R['review'][ri]['tokens']), i['reason'], i['check'], R['review'][ri]['pt'][0], R['review'][ri]['pt'][1], ''] for ri, i in inferred])
    sheet('Lots to review', ['Reason', 'What we need to know', 'H-number candidate', 'Section', 'Lot', 'Label on drawing', 'Map x', 'Map y', 'Outline'], [[(pvm[ri]['reason'] if ri in pvm else ' | '.join(r.get('reasons', []))), (pvm[ri]['check'] if ri in pvm else ''), r.get('htext', ''), r.get('section', ''), r.get('legacy', ''), ''.join(r['tokens']), r['pt'][0], r['pt'][1], r['geom']] for ri, r in open_items])
    sheet('No map position', ['H-number', 'Section', 'Lot', 'Expected x', 'Expected y', 'Confidence', 'Note'], [[u['htext'], u['section'], u['legacy'], (expd.get((u['section'], u['legacy'])) or {}).get('x', ''), (expd.get((u['section'], u['legacy'])) or {}).get('y', ''), (expd.get((u['section'], u['legacy'])) or {}).get('conf') or '', (expd.get((u['section'], u['legacy'])) or {}).get('reason') or u['note'] or 'No label on the drawing resolves to this decoder lot'] for u in unplaced])
    os.makedirs(f'{out}/site/pdf', exist_ok=True); wb.save(f'{out}/site/pdf/Cedar_Grove_H-Map_lots.xlsx')
    print('wrote lots_with_h_numbers.csv (%d verified + %d inferred) and unresolved_locations.csv (%d)' % (len(L), len(inferred), len(open_items) + len(unplaced)))
if __name__ == '__main__': run()
